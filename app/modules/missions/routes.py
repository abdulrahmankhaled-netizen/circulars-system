"""مسارات المهام."""
import re
from datetime import datetime
from flask import (
    Blueprint, render_template, redirect, url_for,
    flash, request, abort,
)
from flask_login import login_required, current_user

from app.extensions import db
from app.core.permissions import require_permission
from app.core.workflow import WorkflowEngine, available_actions
from app.modules.missions.models import Mission
from app.modules.missions.forms import MissionForm
from app.modules.contracts.models import Contract
from app.modules.requests.models import EquipmentRequest

bp = Blueprint("missions", __name__)


# ==================================================================
#  Number Generator
# ==================================================================
def generate_mission_number():
    year = datetime.utcnow().year
    last = Mission.query.filter(
        Mission.mission_number.like(f"MIS-{year}-%")
    ).order_by(Mission.id.desc()).first()
    if last:
        try:
            n = int(last.mission_number.split("-")[-1]) + 1
        except (ValueError, IndexError):
            n = 1
    else:
        n = 1
    return f"MIS-{year}-{n:04d}"


# ==================================================================
#  Equipment Detection
# ==================================================================
def _detect_equipment_key(text):
    """يكتشف نوع المعدة من نص Excel (Actual Rec Type)."""
    if not text:
        return None
    t = str(text).lower().strip()
    t = re.sub(r"\s+", " ", t)

    # 1) مطابقة قاعدة البيانات
    try:
        from app.modules.equipment.models import EquipmentType
        types = EquipmentType.query.filter_by(is_active=True).all()
        for eq in types:
            if eq.key.lower() == t:
                return eq.key
            if eq.name_en and eq.name_en.lower() == t:
                return eq.key
            if eq.name_ar and eq.name_ar == text:
                return eq.key
    except Exception:
        pass

    # 2) مطابقة constants
    from app.core.constants import EQUIPMENT_TYPES
    for eq in EQUIPMENT_TYPES:
        if eq["key"].lower() == t:
            return eq["key"]
        if eq.get("name_en", "").lower() == t:
            return eq["key"]
        if eq.get("name_ar", "") == text:
            return eq["key"]

    # 3) مطابقة كلمات مفتاحية
    rules = [
        (["25", "double"], [], "flatbed_25_double"),
        (["25", "ton"], [], "flatbed_25"),
        (["25t"], [], "flatbed_25"),
        (["13", "double"], [], "flatbed_13_double"),
        (["13", "ton"], [], "flatbed_13"),
        (["13t"], [], "flatbed_13"),
        (["10", "ton"], [], "flatbed_10"),
        (["10t"], [], "flatbed_10"),
        (["7", "ton"], [], "flatbed_7"),
        (["7t"], [], "flatbed_7"),
        (["6", "double"], [], "flatbed_6_double"),
        (["6", "full"], [], "flatbed_6_full"),
        (["6", "half"], [], "flatbed_6_half"),
        (["6", "ton"], [], "flatbed_6"),
        (["6t"], [], "flatbed_6"),
        (["3", "double"], [], "flatbed_3"),
        (["3", "half"], [], "flatbed_3"),
        (["3", "ton"], [], "flatbed_3"),
        (["3t"], [], "flatbed_3"),
        (["mobile", "crane", "100"], [], "mobile_crane_100"),
        (["mobile", "crane", "80"], [], "mobile_crane_80"),
        (["mobile", "crane", "70"], [], "mobile_crane_70"),
        (["mobile", "crane"], [], "mobile_crane_100"),
        (["كرين", "100"], [], "mobile_crane_100"),
        (["كرين", "80"], [], "mobile_crane_80"),
        (["كرين", "70"], [], "mobile_crane_70"),
        (["كرين"], [], "mobile_crane_100"),
        (["wrecker", "heavy"], [], "wrecker_heavy"),
        (["wrecker", "light"], [], "wrecker_light"),
        (["wrecker"], [], "wrecker_heavy"),
        (["ونش", "ثقيل"], [], "wrecker_heavy"),
        (["ونش", "خفيف"], [], "wrecker_light"),
        (["ونش"], [], "wrecker_heavy"),
        (["rucker"], [], "rucker"),
        (["raker"], [], "rucker"),
        (["lowbed"], [], "lowbed"),
        (["low", "bed"], [], "lowbed"),
        (["لوبيد"], [], "lowbed"),
    ]

    for keywords, negative, key in rules:
        if all(kw in t for kw in keywords):
            if not any(neg in t for neg in negative):
                return key

    return None

# ==================================================================
#  Equipment Cost
# ==================================================================
# ==================================================================
#  Equipment Cost
# ==================================================================
def _get_equipment_cost(contract_id, equipment_key, mission_date=None, current_mission_id=None):
    """يجلب سعر المورد من العقد. يوزع التكلفة الشهرية إذا لزم الأمر."""
    from app.modules.contracts.models import Contract, ContractEquipment
    from app.modules.equipment.models import EquipmentPrice
    
    contract = db.session.get(Contract, contract_id)
    if not contract:
        return 0.0

    # 1. البحث في العقد
    contract_equip = ContractEquipment.query.filter_by(
        contract_id=contract_id,
        equipment_key=equipment_key
    ).first()
    
    if contract_equip and contract_equip.unit_price > 0:
        if contract_equip.pricing_period == 'per_mission':
            return float(contract_equip.unit_price)
        elif contract_equip.pricing_period in ['monthly', 'annual']:
            if mission_date:
                # حساب بداية ونهاية الشهر
                month_start = mission_date.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
                if month_start.month == 12:
                    next_month = month_start.replace(year=month_start.year + 1, month=1)
                else:
                    next_month = month_start.replace(month=month_start.month + 1)

                # عدد المهام الحالية في نفس الشهر لنفس المعدة
                query = Mission.query.filter(
                    Mission.contract_id == contract_id,
                    Mission.equipment_key == equipment_key,
                    Mission.mission_date >= month_start,
                    Mission.mission_date < next_month
                )
                
                # إذا كنا نعدل مهمة موجودة، نستثنيها من العدد
                if current_mission_id:
                    query = query.filter(Mission.id != current_mission_id)
                
                count = query.count()
                
                monthly_rate = (contract_equip.quantity or 1) * (contract_equip.unit_price or 0)
                if contract_equip.pricing_period == 'annual':
                    monthly_rate /= 12
                    
                return monthly_rate / (count + 1)
            return 0.0

    # 2. إذا لم يوجد في العقد، نبحث في قائمة أسعار المورد العامة (تُحسب لكل مهمة)
    price = EquipmentPrice.query.filter_by(
        supplier_id=contract.supplier_id,
        equipment_key=equipment_key,
        is_active=True,
    ).first()

    return float(price.unit_price) if price else 0.0# ==================================================================
#  Excel Helpers
# ==================================================================
def _parse_excel_date(value):
    """يحلّل التواريخ من صيغة Excel."""
    if not value:
        return None
    if hasattr(value, "year"):
        return value

    s = str(value).strip()
    formats = [
        "%d-%b-%Y %I:%M:%S %p",
        "%d-%b-%Y %H:%M:%S",
        "%d-%b-%Y %I:%M %p",
        "%d-%b-%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y",
    ]
    for fmt in formats:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def _parse_float(value, default=0.0):
    if value is None or value == "" or value == "N/A":
        return default
    try:
        return float(str(value).replace(",", "").strip())
    except (ValueError, TypeError):
        return default


# ==================================================================
#  List
# ==================================================================
@bp.route("/")
@login_required
@require_permission("missions.view")
def list_missions():
    q = request.args.get("q", "").strip()
    filter_field = request.args.get("filter_field", "").strip()
    filter_value = request.args.get("filter_value", "").strip()
    contract_filter = request.args.get("contract", "").strip() # للتوافق مع الروابط القديمة

    query = Mission.query

    if current_user.role_name == "transport_user":
        query = query.filter(Mission.created_by_id == current_user.id)

    # البحث النصي
    if q:
        like = f"%{q}%"
        query = query.filter(
            db.or_(
                Mission.mission_number.ilike(like),
                Mission.job_order_id.ilike(like),
                Mission.vehicle_plate.ilike(like),
                Mission.customer_name.ilike(like),
                Mission.pickup_location.ilike(like),
                Mission.job_id_external.ilike(like),
            )
        )

    # ⭐ الفلاتر المترابطة
    if filter_field and filter_value:
        if filter_field == "equipment":
            query = query.filter(Mission.equipment_key == filter_value)
        elif filter_field == "status":
            query = query.filter(Mission.status == filter_value)
        elif filter_field == "contract":
            try:
                query = query.filter(Mission.contract_id == int(filter_value))
            except ValueError:
                pass
        elif filter_field == "region":
            query = query.filter(Mission.region == filter_value)
        elif filter_field == "payment":
            query = query.filter(Mission.supplier_payment_status == filter_value)
    
    # للتوافق مع الروابط القديمة
    if contract_filter and not filter_field:
        try:
            query = query.filter(Mission.contract_id == int(contract_filter))
        except ValueError:
            pass

    missions = query.order_by(Mission.mission_date.desc()).all()

    # إحصائيات
      # ⭐ الحسبة العبقرية: حساب الإحصائيات بناءً على المهام المفلترة فقط ⭐
    # (المتغير missions يحتوي فقط على المهام التي ظهرت بعد الفلترة)
    
    total_cost = sum(m.cost or 0 for m in missions)
    total_revenue = sum(m.revenue or 0 for m in missions)
    total_profit = sum(m.profit or 0 for m in missions)

    stats = {
        "total": len(missions), # ⭐ عدد المهام المفلترة
        "draft": sum(1 for m in missions if m.status == "draft"),
        "assigned": sum(1 for m in missions if m.status == "assigned"),
        "in_progress": sum(1 for m in missions if m.status == "in_progress"),
        "completed": sum(1 for m in missions if m.status == "completed"),
        "verified": sum(1 for m in missions if m.status == "verified"),
        "paid": sum(1 for m in missions if m.status == "paid"),
        "total_cost": float(total_cost),
        "total_revenue": float(total_revenue),
        "total_profit": float(total_profit),
    }

    contracts = Contract.query.order_by(Contract.contract_number).all()

    # ⭐ تجهيز قوائم الفلاتر المترابطة (من قاعدة البيانات)
    from app.core.constants import MISSION_STATES, PAYMENT_STATUSES
    from app.modules.equipment.models import EquipmentType
    
    # 1. المعدات (من جدول EquipmentType)
    equipment_options = []
    try:
        types = EquipmentType.query.filter_by(is_active=True)\
            .order_by(EquipmentType.sort_order, EquipmentType.name_ar).all()
        equipment_options = [
            {"val": t.key, "text": f"{t.name_ar} ({t.name_en or ''})".strip()} 
            for t in types
        ]
    except Exception as e:
        print(f"Error loading equipment types: {e}")
    
    # 2. حالات المهمة
    status_options = [{"val": s[0], "text": s[1]} for s in MISSION_STATES]
    
    # 3. العقود
    contract_options = [{"val": str(c.id), "text": c.contract_number} for c in contracts]
    
    # 4. المناطق (استخراج المناطق الفريدة من المهام)
    unique_regions = db.session.query(Mission.region).filter(Mission.region.isnot(None)).distinct().all()
    region_options = [{"val": r[0], "text": r[0]} for r in unique_regions]
    
    # 5. حالات الدفع
    payment_options = [{"val": p[0], "text": p[1]} for p in PAYMENT_STATUSES]

    filter_options = {
        "equipment": equipment_options,
        "status": status_options,
        "contract": contract_options,
        "region": region_options,
        "payment": payment_options,
    }

    return render_template(
        "missions/list.html",
        missions=missions,
        contracts=contracts,
        q=q,
        filter_field=filter_field,
        filter_value=filter_value,
        filter_options=filter_options,
        stats=stats,
    )# ==================================================================
#  Add
# ==================================================================
@bp.route("/add", methods=["GET", "POST"])
@login_required
@require_permission("missions.create")
def add_mission():
    form = MissionForm()
    form.contract_id.choices = [
        (c.id, f"{c.contract_number} - {c.supplier.company_name_ar} [{c.status}]")
        for c in Contract.query.order_by(Contract.contract_number).all()
    ]

    if request.method == "GET":
        cid = request.args.get("contract", type=int)
        if cid:
            form.contract_id.data = cid
        rid = request.args.get("request", type=int)
        if rid:
            req = db.session.get(EquipmentRequest, rid)
            if req:
                form.customer_name.data = req.client_entity
                form.region.data = req.region
                form.pickup_location.data = req.location

    if form.validate_on_submit():
        try:
            auto_cost = _get_equipment_cost(
                form.contract_id.data,
                form.equipment_key.data,
            )
            final_cost = form.cost.data if (form.cost.data and form.cost.data > 0) else auto_cost

            m = Mission(
                mission_number=generate_mission_number(),
                contract_id=form.contract_id.data,
                equipment_key=form.equipment_key.data,
                vehicle_plate=form.vehicle_plate.data,
                vehicle_make=form.vehicle_make.data,
                vehicle_model=form.vehicle_model.data,
                vehicle_vin=form.vehicle_vin.data,
                customer_name=form.customer_name.data,
                customer_phone=form.customer_phone.data,
                pickup_location=form.pickup_location.data,
                drop_location=form.drop_location.data,
                region=form.region.data,
                mission_date=form.mission_date.data or datetime.utcnow(),
                driver_name=form.driver_name.data,
                driver_phone=form.driver_phone.data,
                cost=final_cost or 0.0,
                revenue=form.revenue.data or 0.0,
                supplier_payment_status=form.supplier_payment_status.data,
                customer_invoice_status=form.customer_invoice_status.data,
                status=form.status.data,
                notes=form.notes.data,
                completion_notes=form.completion_notes.data,
                created_by_id=current_user.id,
                is_impound=form.is_impound.data,
                store_number=form.store_number.data,
                job_order_id=form.job_order_id.data,
                job_id_external=form.job_id_external.data,
                payer_type=form.payer_type.data,
                requested_rec_type=form.requested_rec_type.data,
                actual_rec_type=form.actual_rec_type.data,
                toll_job_id=form.toll_job_id.data,
                toll_amount=form.toll_amount.data,
                advisor=form.advisor.data,
                service_centre=form.service_centre.data,
                sub_status=form.sub_status.data,
            )
            m.recalculate_financials()

            rid = request.args.get("request", type=int)
            if rid:
                m.request_id = rid

            db.session.add(m)
            db.session.commit()
            flash(f"تم إنشاء المهمة {m.mission_number}", "success")
            return redirect(url_for("missions.view_mission", mid=m.id))

        except Exception as e:
            db.session.rollback()
            import traceback
            print(traceback.format_exc())
            flash(f"خطأ في الحفظ: {str(e)[:200]}", "danger")

    return render_template(
        "missions/form.html",
        form=form,
        title="مهمة جديدة",
        mission=None,
    )


# ==================================================================
#  Edit
# ==================================================================
@bp.route("/<int:mid>/edit", methods=["GET", "POST"])
@login_required
@require_permission("missions.edit")
def edit_mission(mid):
    m = db.session.get(Mission, mid)
    if not m:
        abort(404)

    form = MissionForm(obj=m)
    form.contract_id.choices = [
        (c.id, f"{c.contract_number} - {c.supplier.company_name_ar} [{c.status}]")
        for c in Contract.query.order_by(Contract.contract_number).all()
    ]

    if form.validate_on_submit():
        m.contract_id = form.contract_id.data
        m.equipment_key = form.equipment_key.data
        m.vehicle_plate = form.vehicle_plate.data
        m.vehicle_make = form.vehicle_make.data
        m.vehicle_model = form.vehicle_model.data
        m.vehicle_vin = form.vehicle_vin.data
        m.customer_name = form.customer_name.data
        m.customer_phone = form.customer_phone.data
        m.pickup_location = form.pickup_location.data
        m.drop_location = form.drop_location.data
        m.region = form.region.data
        m.mission_date = form.mission_date.data or m.mission_date
        m.driver_name = form.driver_name.data
        m.driver_phone = form.driver_phone.data
        m.revenue = form.revenue.data or 0.0
        m.supplier_payment_status = form.supplier_payment_status.data
        m.customer_invoice_status = form.customer_invoice_status.data
        m.status = form.status.data
        m.notes = form.notes.data
        m.completion_notes = form.completion_notes.data

        m.is_impound = form.is_impound.data
        m.store_number = form.store_number.data
        m.job_order_id = form.job_order_id.data

        m.job_id_external = form.job_id_external.data
        m.payer_type = form.payer_type.data
        m.requested_rec_type = form.requested_rec_type.data
        m.actual_rec_type = form.actual_rec_type.data
        m.toll_job_id = form.toll_job_id.data
        m.toll_amount = form.toll_amount.data
        m.advisor = form.advisor.data
        m.service_centre = form.service_centre.data
        m.sub_status = form.sub_status.data

        if form.cost.data and form.cost.data > 0:
            m.cost = form.cost.data
        else:
            auto_cost = _get_equipment_cost(m.contract_id, form.equipment_key.data)
            m.cost = auto_cost if auto_cost > 0 else (m.cost or 0.0)

        m.recalculate_financials()
        db.session.commit()
        flash("تم تحديث المهمة", "success")
        return redirect(url_for("missions.view_mission", mid=m.id))

    return render_template(
        "missions/form.html",
        form=form,
        title=f"تعديل: {m.mission_number}",
        mission=m,
    )


# ==================================================================
#  View
# ==================================================================
@bp.route("/<int:mid>")
@login_required
@require_permission("missions.view")
def view_mission(mid):
    m = db.session.get(Mission, mid)
    if not m:
        abort(404)

    actions = available_actions("mission", m.status, current_user)

    from app.modules.workflow.models import WorkflowLog
    logs = WorkflowLog.query.filter_by(
        entity_type="mission", entity_id=m.id
    ).order_by(WorkflowLog.created_at.asc()).all()

    return render_template(
        "missions/view.html",
        mission=m,
        actions=actions,
        logs=logs,
    )


# ==================================================================
#  Workflow Actions
# ==================================================================
@bp.route("/<int:mid>/action", methods=["POST"])
@login_required
@require_permission("missions.edit")
def action_mission(mid):
    m = db.session.get(Mission, mid)
    if not m:
        abort(404)

    action = request.form.get("action", "").strip()
    comment = request.form.get("comment", "").strip()

    if not action:
        flash("لم يتم تحديد إجراء", "danger")
        return redirect(url_for("missions.view_mission", mid=mid))

    engine = WorkflowEngine("mission")
    ok, err = engine.execute(m, action, user=current_user, comment=comment)

    if ok:
        if action == "assign":
            m.assigned_to_id = current_user.id
            m.assigned_at = datetime.utcnow()
        elif action == "start":
            m.started_at = datetime.utcnow()
        elif action == "complete":
            m.completed_at = datetime.utcnow()
        elif action == "verify":
            m.verified_by_id = current_user.id
            m.verified_at = datetime.utcnow()
        elif action == "invoiced":
            m.customer_invoice_status = "invoiced"
        elif action == "paid":
            m.supplier_payment_status = "paid"
            m.customer_invoice_status = "paid"

        db.session.commit()
        flash(f"تم تنفيذ '{action}'", "success")
    else:
        flash(f"{err}", "danger")

    return redirect(url_for("missions.view_mission", mid=mid))


# ==================================================================
#  Delete
# ==================================================================
@bp.route("/<int:mid>/delete", methods=["POST"])
@login_required
@require_permission("missions.delete")
def delete_mission(mid):
    m = db.session.get(Mission, mid)
    if not m:
        abort(404)

    if not current_user.is_admin and m.status in ("paid", "invoiced"):
        flash("لا يمكن حذف مهمة مدفوعة أو مفوترة", "warning")
        return redirect(url_for("missions.view_mission", mid=mid))

    number = m.mission_number
    db.session.delete(m)
    db.session.commit()
    flash(f"تم حذف المهمة {number}", "warning")
    return redirect(url_for("missions.list_missions"))


# ==================================================================
#  Bulk Delete
# ==================================================================
@bp.route("/bulk-delete", methods=["POST"])
@login_required
@require_permission("missions.delete")
def bulk_delete():
    ids = request.form.getlist("mission_ids")

    if not ids:
        flash("لم يتم تحديد أي مهمة", "warning")
        return redirect(url_for("missions.list_missions"))

    if not current_user.is_admin:
        flash("الحذف الجماعي متاح لمدير النظام فقط", "danger")
        return redirect(url_for("missions.list_missions"))

    deleted = 0
    skipped = 0

    for mid in ids:
        try:
            m = db.session.get(Mission, int(mid))
            if not m:
                skipped += 1
                continue
            if m.status in ("paid", "invoiced"):
                skipped += 1
                continue
            db.session.delete(m)
            deleted += 1
        except (ValueError, TypeError):
            skipped += 1

    db.session.commit()

    msg = f"تم حذف {deleted} مهمة"
    if skipped > 0:
        msg += f" - تم تخطي {skipped}"
    flash(msg, "success" if deleted > 0 else "warning")

    return redirect(url_for("missions.list_missions"))


# ==================================================================
#  Recalculate Costs
# ==================================================================
@bp.route("/recalculate-costs", methods=["POST"])
@login_required
@require_permission("missions.edit")
def recalculate_costs():
    if not current_user.is_admin:
        flash("متاح لمدير النظام فقط", "danger")
        return redirect(url_for("missions.list_missions"))

    from app.modules.contracts.models import ContractEquipment
    from app.modules.equipment.models import EquipmentPrice
    from collections import defaultdict

    missions = Mission.query.all()
    updated = 0
    groups = defaultdict(list)
    
    for m in missions:
        contract = m.contract
        if not contract:
            continue

        # 1. البحث عن المعدة بالمفتاح (Key) أولاً
        contract_equip = ContractEquipment.query.filter_by(
            contract_id=m.contract_id, 
            equipment_key=m.equipment_key
        ).first()

        # 2. إذا لم يجدها، يبحث بالاسم العربي أو الإنجليزي
        if not contract_equip:
            m_name_ar = m.equipment_name_ar()
            m_name_en = m.equipment_name_en()
            for ce in contract.equipment_items:
                if (ce.equipment_name_ar() == m_name_ar) or (m_name_en and ce.equipment_name_en() == m_name_en):
                    contract_equip = ce
                    break

        # 3. الحساب بناءً على ما وجده
        if contract_equip:
            if contract_equip.pricing_period == 'per_mission':
                m.cost = contract_equip.unit_price
                m.recalculate_financials()
                updated += 1
            else:
                if m.mission_date:
                    month_key = m.mission_date.strftime('%Y-%m')
                    groups[(m.contract_id, contract_equip.equipment_key, month_key)].append(m)
        else:
            # البحث في قائمة أسعار المورد العامة
            price = EquipmentPrice.query.filter_by(
                supplier_id=contract.supplier_id, 
                equipment_key=m.equipment_key, 
                is_active=True
            ).first()
            if price and price.unit_price > 0:
                m.cost = price.unit_price
                m.recalculate_financials()
                updated += 1

    # توزيع التكاليف الشهرية (إذا وجدت)
    for (cid, eq_key, month_key), group_missions in groups.items():
        contract_equip = ContractEquipment.query.filter_by(
            contract_id=cid, 
            equipment_key=eq_key
        ).first()
        
        if contract_equip and contract_equip.pricing_period in ['monthly', 'annual']:
            monthly_cost = (contract_equip.quantity or 1) * (contract_equip.unit_price or 0)
            if contract_equip.pricing_period == 'annual':
                monthly_cost /= 12
                
            cost_per_mission = monthly_cost / len(group_missions)
            
            for m in group_missions:
                m.cost = cost_per_mission
                m.recalculate_financials()
                updated += 1

    db.session.commit()
    flash(f"✅ تم إعادة حساب {updated} مهمة بنجاح. (تم مطابقة الأسماء بذكاء).", "success")
    return redirect(url_for("missions.list_missions"))# ==================================================================
#  Excel Upload
# ==================================================================
@bp.route("/upload", methods=["GET", "POST"])
@login_required
@require_permission("missions.create")
def upload_excel():
    contracts = Contract.query.order_by(Contract.status.desc(), Contract.project_name, Contract.contract_number).all()

    if request.method == "POST":
        file = request.files.get("excel_file")
        if not file or not file.filename:
            flash("الرجاء اختيار ملف Excel", "warning")
            return redirect(url_for("missions.upload_excel"))

        if not file.filename.lower().endswith((".xlsx", ".xlsm", ".xls")):
            flash("نوع الملف يجب أن يكون Excel", "danger")
            return redirect(url_for("missions.upload_excel"))

        default_contract_id = request.form.get("default_contract_id", "").strip()
        default_cid = None
        try:
            default_cid = int(default_contract_id) if default_contract_id else None
        except ValueError:
            pass

        if not default_cid:
            flash("⚠️ يجب اختيار العقد أولاً", "warning")
            return redirect(url_for("missions.upload_excel"))

        default_contract = db.session.get(Contract, default_cid)
        contract_equips = default_contract.equipment_items if default_contract else []

        try:
            from openpyxl import load_workbook
            wb = load_workbook(file, data_only=True)
            ws = wb.active

            headers_row = None
            for r in range(1, 6):
                row_vals = [cell.value for cell in ws[r]]
                if any(v and "Job Id" in str(v) for v in row_vals):
                    headers_row = r
                    break

            if headers_row is None:
                flash("❌ لم يُعثر على صف الرؤوس (يجب أن يحتوي على 'Job Id')", "danger")
                return redirect(url_for("missions.upload_excel"))

            headers = []
            for cell in ws[headers_row]:
                v = cell.value
                headers.append(str(v).strip() if v else "")

            def col(name):
                try:
                    return headers.index(name)
                except ValueError:
                    return None

            col_map = {
                "job_id":          col("Job Id"),
                "vin":             col("VIN"),
                "plate":           col("Plate"),
                "payer_type":      col("Payer Type"),
                "created_date":    col("Created Date"),
                "started_date":    col("Started Date"),
                "completed_date":  col("Completed Date"),
                "pickup_location": col("Pickup Location"),
                "drop_location":   col("Drop Location"),
                "status":          col("Status"),
                "driver_name":     col("Driver Name"),
                "driver_phone":    col("Driver Phone"),
                "amount":          col("Amount"),
                "requested_rec":   col("Rquested Rec Type"),
                "actual_rec":      col("Actual Rec Type"),
                "make":            col("Make"),
                "model":           col("Model"),
                "toll_job_id":     col("Toll Job Id"),
                "toll_amount":     col("Toll Amount"),
                "remarks":         col("Remarks"),
                "sub_status":      col("Sub Status"),
                "advisor":         col("Advisor"),
                "service_centre":  col("Service Centre/ Workshop"),
                "store_number":    col("Store Number"),
                "job_order_id":    col("Job Order"),
                "qr_code":         col("QrCode") or col("QR Code") or col("QRCode"),
            }

            created = 0
            updated = 0
            skipped = 0
            detected = {}

            for row_idx, row in enumerate(
                ws.iter_rows(min_row=headers_row + 1, values_only=True),
                start=headers_row + 1
            ):
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue

                def g(key):
                    idx = col_map.get(key)
                    if idx is None or idx >= len(row):
                        return None
                    v = row[idx]
                    if v is None or str(v).strip() == "" or str(v).strip() == "N/A":
                        return None
                    return v

                job_id = g("job_id")
                if not job_id:
                    skipped += 1
                    continue

                job_id_str = str(job_id).strip()

                existing = Mission.query.filter(
                    db.or_(
                        Mission.job_order_id == job_id_str,
                        Mission.job_id_external == job_id_str,
                        Mission.mission_number == job_id_str
                    )
                ).first()

                created_dt = _parse_excel_date(g("created_date"))
                started_dt = _parse_excel_date(g("started_date"))
                completed_dt = _parse_excel_date(g("completed_date"))

                actual_rec = g("actual_rec")
                requested_rec = g("requested_rec")
                rec_type = actual_rec or requested_rec
                
                equipment_key = None
                rec_type_str = str(rec_type or "").strip().lower()
                
                if contract_equips and rec_type_str:
                    for c_eq in contract_equips:
                        eq_name_en = (c_eq.equipment_name_en() or "").lower()
                        eq_name_ar = (c_eq.equipment_name_ar() or "").lower()
                        
                        if eq_name_en and eq_name_en in rec_type_str:
                            equipment_key = c_eq.equipment_key
                            break
                        if eq_name_ar and eq_name_ar in rec_type_str:
                            equipment_key = c_eq.equipment_key
                            break
                        if "mobile crane" in rec_type_str and "mobile crane" in eq_name_en:
                            equipment_key = c_eq.equipment_key
                            break

                if not equipment_key:
                    equipment_key = _detect_equipment_key(rec_type) or "flatbed_6"

                detected[str(rec_type)] = equipment_key

                amount = _parse_float(g("amount"), 0.0)
                excel_status = str(g("status") or "").strip()
                mission_status = "completed" if excel_status == "Job Completed" else "assigned"

                if existing:
                    m = existing
                else:
                    m = Mission(
                        mission_number=generate_mission_number(),
                        contract_id=default_cid,
                        created_by_id=current_user.id,
                    )
                    db.session.add(m)

                m.job_id_external = job_id_str
                m.job_order_id = job_id_str 
                m.contract_id = default_cid
                m.vehicle_vin = str(g("vin") or "").strip() or None
                m.vehicle_plate = str(g("plate") or "").strip()[:40] or None
                m.payer_type = str(g("payer_type") or "").strip()[:50] or None
                m.mission_date = created_dt or datetime.utcnow()
                m.created_date = created_dt
                m.started_date = started_dt
                m.completed_date = completed_dt
                m.started_at = started_dt
                m.completed_at = completed_dt
                m.pickup_location = str(g("pickup_location") or "").strip()[:255] or None
                m.drop_location = str(g("drop_location") or "").strip()[:255] or None
                pickup = g("pickup_location")
                m.region = str(pickup).split()[0] if pickup else None
                m.job_status = excel_status[:50] or None
                m.driver_name = str(g("driver_name") or "").strip()[:120] or None
                m.driver_phone = str(g("driver_phone") or "").strip()[:40] or None
                m.equipment_key = equipment_key
                m.requested_rec_type = str(requested_rec or "").strip()[:120] or None
                m.actual_rec_type = str(actual_rec or "").strip()[:120] or None
                m.vehicle_make = str(g("make") or "").strip()[:80] or None
                m.vehicle_model = str(g("model") or "").strip()[:80] or None
                m.toll_job_id = str(g("toll_job_id") or "").strip()[:80] or None
                m.toll_amount = _parse_float(g("toll_amount"), None)
                m.notes = str(g("remarks") or "").strip() or None
                m.sub_status = str(g("sub_status") or "").strip()[:80] or None
                m.advisor = str(g("advisor") or "").strip()[:120] or None
                m.service_centre = str(g("service_centre") or "").strip()[:150] or None

                store_val = g("store_number")
                if store_val:
                    m.store_number = str(store_val).strip()[:80]
                    m.is_impound = True

                jo_val = g("job_order_id")
                if jo_val:
                    m.job_order_id = str(jo_val).strip()[:80]

                qr_val = g("qr_code")
                if qr_val:
                    m.qr_code = str(qr_val).strip()

                m.revenue = amount
                
                cost_from_price = _get_equipment_cost(m.contract_id, equipment_key, m.mission_date)
                
                if cost_from_price > 0:
                    m.cost = cost_from_price
                elif not m.cost:
                    m.cost = 0.0

                m.recalculate_financials()
                m.status = mission_status

                if started_dt:
                    m.mission_date = started_dt

                if existing:
                    updated += 1
                else:
                    created += 1

            db.session.commit()

            msg = f"✅ استيراد: {created} جديد، {updated} محدّث، {skipped} متخطى"
            flash(msg, "success")

            if detected:
                summary = " | ".join(
                    f"'{k[:25]}' -> {v}"
                    for k, v in list(detected.items())[:5]
                )
                flash(f"🔍 الأصناف المكتشفة: {summary}", "info")

            return redirect(url_for("missions.list_missions"))

        except Exception as e:
            db.session.rollback()
            import traceback
            print(traceback.format_exc())
            flash(f"❌ خطأ في قراءة Excel: {e}", "danger")
            return redirect(url_for("missions.upload_excel"))

    return render_template(
        "missions/upload.html",
        contracts=contracts,
    )
# ==================================================================
#  Download Template
# ==================================================================
@bp.route("/download-template")
@login_required
@require_permission("missions.create")
def download_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.sheet_view.rightToLeft = False

    ws["A1"] = "Project Jobs List - Fleet Management"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    ws.merge_cells("A1:Z1")
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = [
        "Job Id", "VIN", "Plate", "Payer Type", "Created Date", "Started Date",
        "Completed Date", "Pickup Location", "Drop Location", "Status",
        "Driver Name", "Driver Phone", "Amount", "QrCode",
        "Rquested Rec Type", "Actual Rec Type", "Make", "Model",
        "Toll Job Id", "Toll Amount", "Remarks", "Sub Status",
        "Imp. Form No", "Test Job", "Advisor", "Service Centre/ Workshop",
    ]

    hfont = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill("solid", fgColor="1F4E78")
    center = Alignment(horizontal="center", vertical="center")

    for idx, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=idx, value=h)
        c.font = hfont
        c.fill = hfill
        c.alignment = center

    sample = [
        2823681, "N/A", "Dubai Public Transportation 36517", "Cash Customer",
        "01-Oct-2026 04:09:09 PM", "01-Oct-2026 04:09:47 PM", "01-Oct-2026 04:12:16 PM",
        "Fujairah", "Fujairah", "Job Completed",
        "MUHAMMAD QAYUM KHAN NABI RAHAM (1029)", "0549903669", 3150, "2823681",
        "Wrecker heavy", "6 TON Double Towing", "Volvo", "FH12",
        "", "", "", "", "", "No", "", "",
    ]
    for idx, v in enumerate(sample, 1):
        ws.cell(row=3, column=idx, value=v)

    widths = [12, 10, 30, 15, 22, 22, 22, 25, 25, 15,
              30, 15, 10, 15, 20, 20, 14, 14,
              12, 12, 20, 14, 12, 10, 14, 25]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    from flask import send_file
    return send_file(
        bio,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"missions_template_{datetime.utcnow().strftime('%Y%m%d')}.xlsx",
    )
# ==================================================================
#  Payment Follow-up (متابعة الدفع)
# ==================================================================
@bp.route("/followup")
@login_required
@require_permission("missions.view")
def followup_list():
    """صفحة متابعة المركبات المعلقة."""
    status_filter = request.args.get("status", "").strip()
    
    query = Mission.query
    if status_filter:
        query = query.filter(Mission.followup_status == status_filter)
    
    missions = query.order_by(Mission.mission_date.desc()).all()

    # إحصائيات البابلز
    stats = {
        "pending": Mission.query.filter_by(followup_status="pending").count(),
        "circular": Mission.query.filter_by(followup_status="circular").count(),
        "paid": Mission.query.filter_by(followup_status="paid").count(),
    }

    return render_template("missions/followup.html", missions=missions, stats=stats, status_filter=status_filter)


@bp.route("/followup/<int:mid>/update", methods=["POST"])
@login_required
@require_permission("missions.edit")
def followup_update(mid):
    """تحديث حالة المتابعة للمركبة."""
    m = db.session.get(Mission, mid)
    if not m:
        abort(404)
    
    new_status = request.form.get("followup_status")
    if new_status in ["pending", "circular", "paid"]:
        m.followup_status = new_status
        db.session.commit()
        flash(f"✅ تم تحديث حالة المهمة {m.mission_number} إلى {new_status}", "success")
    
    return redirect(url_for("missions.followup_list"))
# ==================================================================
#  تحديث جماعي (تغيير حالة المهام إلى مدفوعة)
# ==================================================================
@bp.route("/bulk-mark-paid", methods=["POST"])
@login_required
@require_permission("missions.edit")
def bulk_mark_paid():
    # 1. استلام النص الذي قمت بلصقه
    text = request.form.get("mission_ids_text", "").strip()
    if not text:
        flash("⚠️ الرجاء لصق أرقام المهام أولاً", "warning")
        return redirect(url_for("missions.list_missions"))

    # 2. فصل الأرقام عن بعضها (سواء كانت مفصولة بمسافة أو سطر جديد)
    import re
    raw_ids = re.split(r'[\n,\t\s]+', text)
    search_ids = [i.strip() for i in raw_ids if i.strip()]

    if not search_ids:
        flash("⚠️ لم يتم العثور على أرقام صالحة", "warning")
        return redirect(url_for("missions.list_missions"))

    # 3. البحث في قاعدة البيانات عن هذه الأرقام
    missions = Mission.query.filter(
        db.or_(
            Mission.mission_number.in_(search_ids),
            Mission.job_order_id.in_(search_ids),
            Mission.job_id_external.in_(search_ids)
        )
    ).all()

    if not missions:
        flash("❌ لم يتم العثور على أي مهام مطابقة للأرقام المدخلة.", "danger")
        return redirect(url_for("missions.upload_excel"))

    # 4. تغيير حالتها إلى مدفوعة
    updated_count = 0
    for m in missions:
        m.supplier_payment_status = "paid"
        m.customer_invoice_status = "paid"
        m.followup_status = "paid"
        updated_count += 1

    db.session.commit()

    flash(f"✅ تم تحديث حالة {updated_count} مهمة إلى 'مدفوعة' بنجاح.", "success")
    return redirect(url_for("missions.list_missions"))
# ==================================================================
#  تصدير متابعة الدفع إلى Excel
# ==================================================================
# ==================================================================
#  تصدير متابعة الدفع إلى Excel (مُحدث ليشمل السعر)
# ==================================================================
@bp.route("/followup/export")
@login_required
@require_permission("missions.view")
def export_followup_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from flask import send_file
    import io

    status_filter = request.args.get("status", "").strip()
    query = Mission.query
    if status_filter:
        query = query.filter(Mission.followup_status == status_filter)
    missions = query.order_by(Mission.mission_date.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "متابعة الدفع"
    ws.sheet_view.rightToLeft = True

    # ⭐ أضفنا عمود "المبلغ المطلوب من العميل"
    headers = [
        "رقم المهمة", 
        "Job Order", 
        "اللوحة", 
        "المعدة", 
        "التاريخ", 
        "المبلغ المطلوب من العميل (AED)", 
        "حالة الدفع (اختر من القائمة)"
    ]
    
    hfont = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1F4E78")
    center = Alignment(horizontal="center")

    for idx, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=idx, value=h)
        c.font = hfont
        c.fill = hfill
        c.alignment = center

    # تحويل الحالات إلى نصوص عربية
    status_labels = {
        "pending": "معلق بانتظار الدفع", 
        "circular": "تم رفع تعميم", 
        "paid": "تم الدفع"
    }

    for i, m in enumerate(missions, 2):
        ws.cell(row=i, column=1, value=m.mission_number)
        ws.cell(row=i, column=2, value=m.job_order_id or m.job_id_external or "")
        ws.cell(row=i, column=3, value=m.vehicle_plate or "")
        ws.cell(row=i, column=4, value=m.equipment_name_ar())
        ws.cell(row=i, column=5, value=m.mission_date.strftime("%Y-%m-%d") if m.mission_date else "")
        
        # ⭐ إضافة قيمة الإيراد (المبلغ المطلوب من العميل)
        ws.cell(row=i, column=6, value=m.revenue or 0)
        
        ws.cell(row=i, column=7, value=status_labels.get(m.followup_status, m.followup_status))

    # إضافة قائمة منسدلة في الإكسل لحقل الحالة (العمود G الآن بدلاً من F)
    dv = DataValidation(type="list", formula1='"معلق بانتظار الدفع,تم رفع تعميم,تم الدفع"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{len(missions)+1}")

    # ضبط عرض الأعمدة لتكون واضحة
    widths = [20, 20, 15, 25, 15, 25, 25]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    return send_file(
        bio,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="followup_missions.xlsx"
    )
# ==================================================================
#  استيراد وتحديث متابعة الدفع من Excel
# ==================================================================
@bp.route("/followup/import", methods=["POST"])
@login_required
@require_permission("missions.edit")
def import_followup_excel():
    file = request.files.get("excel_file")
    if not file or not file.filename:
        flash("⚠️ الرجاء اختيار ملف Excel أولاً", "warning")
        return redirect(url_for("missions.followup_list"))

    try:
        from openpyxl import load_workbook
        wb = load_workbook(file, data_only=True)
        ws = wb.active

        # قراءة العناوين لمعرفة ترتيب الأعمدة
        headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
        col_mission = None
        col_job = None
        col_status = None

        for idx, h in enumerate(headers):
            if "رقم المهمة" in h: col_mission = idx
            if "Job Order" in h: col_job = idx
            if "حالة الدفع" in h: col_status = idx

        if col_status is None:
            flash("❌ لم يتم العثور على عمود (حالة الدفع) في الملف.", "danger")
            return redirect(url_for("missions.followup_list"))

        updated = 0
        # تحويل النص العربي إلى الحالة في قاعدة البيانات
        reverse_labels = {
            "معلق بانتظار الدفع": "pending", "معلق": "pending",
            "تم رفع تعميم": "circular", "رفع تعميم": "circular",
            "تم الدفع": "paid", "مدفوع": "paid"
        }

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row): continue
            
            status_val = str(row[col_status]).strip() if col_status is not None and col_status < len(row) else ""
            if not status_val: continue

            new_status = reverse_labels.get(status_val)
            if not new_status: continue

            # البحث عن المهمة (سواء برقمها أو رقم أمر العمل)
            m = None
            if col_mission is not None and col_mission < len(row) and row[col_mission]:
                m = Mission.query.filter_by(mission_number=str(row[col_mission]).strip()).first()
            
            if not m and col_job is not None and col_job < len(row) and row[col_job]:
                m = Mission.query.filter_by(job_order_id=str(row[col_job]).strip()).first()

            if m:
                m.followup_status = new_status
                updated += 1
        
        db.session.commit()
        flash(f"✅ تم تحديث حالة {updated} مهمة بنجاح من الملف!", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"❌ حدث خطأ أثناء قراءة الملف: {e}", "danger")

    return redirect(url_for("missions.followup_list"))
